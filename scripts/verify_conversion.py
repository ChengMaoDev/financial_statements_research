"""Check content retention and source order for structure-aware derivatives.

This validates extraction integrity, not factual correctness or perfect visual fidelity.
Run with the same Python packages as convert_to_text.py; PDF checks require pdftotext.
"""
from __future__ import annotations

import argparse
import csv
import re
import subprocess
import sys
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "origin"
ERRORS: list[str] = []
CHECKED = 0


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def error(message: str) -> None:
    ERRORS.append(message)
    print("MISMATCH:", message, flush=True)


def serialized_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def check_pdf(source: Path, output: Path) -> None:
    global CHECKED
    CHECKED += 1
    if not output.exists():
        error(f"{output} missing")
        return
    extracted = subprocess.run(
        ["pdftotext", "-layout", str(source), "-"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout
    source_pages = extracted.split("\f")
    if source_pages and not source_pages[-1].strip():
        source_pages.pop()
    text = output.read_text(encoding="utf-8")
    pages = re.findall(r"^## PDF page (\d+)\n\n```text\n(.*?)\n```$", text, re.M | re.S)
    if [int(page_number) for page_number, _ in pages] != list(range(1, len(source_pages) + 1)):
        error(f"{source.name}: page markers do not match {len(source_pages)} source pages")
        return
    for number, ((_, actual), expected) in enumerate(zip(pages, source_pages), 1):
        expected_lines = [line.rstrip() for line in expected.splitlines()]
        while expected_lines and not expected_lines[0].strip():
            expected_lines.pop(0)
        while expected_lines and not expected_lines[-1].strip():
            expected_lines.pop()
        if normalize(actual) != normalize("\n".join(expected_lines)):
            error(f"{source.name} PDF page {number}: extracted text differs")
            return
    print(f"OK PDF {source.name}: {len(source_pages)} page texts retained")


def docx_elements(source: Path) -> Iterable[str]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(source)
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            value = Paragraph(child, document).text.strip()
            if value:
                yield value
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            for row in table.rows:
                for cell in row.cells:
                    value = cell.text.strip()
                    if value:
                        yield value


def check_docx(source: Path, output: Path) -> None:
    global CHECKED
    CHECKED += 1
    if not output.exists():
        error(f"{output} missing")
        return
    rendered = normalize(output.read_text(encoding="utf-8").replace("\\|", "|"))
    cursor = 0
    for index, source_text in enumerate(docx_elements(source), 1):
        target = normalize(source_text)
        position = rendered.find(target, cursor)
        if position < 0:
            error(f"{source.name}: source element {index} missing or out of order: {target[:70]!r}")
            return
        cursor = position + len(target)
    print(f"OK DOCX {source.name}: paragraph/table text retained in source order")


def expected_xlsx_rows(source: Path, sheet_name: str) -> list[list[str]]:
    import openpyxl

    values = openpyxl.load_workbook(source, data_only=True, read_only=False)
    formulas = openpyxl.load_workbook(source, data_only=False, read_only=False)
    expected: list[list[str]] = []
    for value_row, formula_row in zip(values[sheet_name].iter_rows(), formulas[sheet_name].iter_rows()):
        row: list[str] = []
        for value_cell, formula_cell in zip(value_row, formula_row):
            formula = formula_cell.value
            value = value_cell.value
            if isinstance(formula, str) and formula.startswith("="):
                cached = serialized_value(value)
                row.append(
                    f"FORMULA (cached result unavailable): {formula}"
                    if value is None
                    else f"FORMULA: {formula} | CACHED VALUE: {cached}"
                )
            else:
                row.append(serialized_value(value))
        if any(cell.strip() for cell in row):
            expected.append(row)
    values.close()
    formulas.close()
    return expected


def safe_sheet_name(name: str) -> str:
    return re.sub(r"[^\w\-\u4e00-\u9fff]+", "_", name).strip("_") or "sheet"


def check_xlsx(source: Path, outdir: Path) -> None:
    global CHECKED
    import openpyxl

    CHECKED += 1
    workbook = openpyxl.load_workbook(source, data_only=False, read_only=True)
    for sheet in workbook.worksheets:
        csv_path = outdir / f"{safe_sheet_name(sheet.title)}.csv"
        if not csv_path.exists():
            error(f"{source.name}[{sheet.title}]: CSV missing")
            continue
        with csv_path.open(encoding="utf-8-sig", newline="") as stream:
            actual = list(csv.reader(stream))
        expected = expected_xlsx_rows(source, sheet.title)
        if actual != expected:
            mismatch = next(
                ((r, c) for r, (actual_row, expected_row) in enumerate(zip(actual, expected), 1)
                 for c, (a, e) in enumerate(zip(actual_row, expected_row), 1) if a != e),
                None,
            )
            error(f"{source.name}[{sheet.title}]: CSV cell mismatch near {mismatch}; rows {len(actual)} vs {len(expected)}")
            continue
        formulas = sum(
            isinstance(cell.value, str) and cell.value.startswith("=")
            for row in sheet.iter_rows() for cell in row
        )
        print(f"OK XLSX {source.name}[{sheet.title}]: {len(actual)} rows, {formulas} formulas retained")
    workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT, help="Output root (default: repository root)")
    parser.add_argument("--only", choices=["all", "pdf", "docx", "xlsx"], default="all")
    args = parser.parse_args()
    output = args.output_dir.resolve()
    pdfs = sorted((SRC / "年度报告").glob("*.pdf")) if args.only in ("all", "pdf") else []
    docxs = sorted(SRC.glob("*.docx")) if args.only in ("all", "docx") else []
    xlsxs = sorted(SRC.glob("*.xlsx")) if args.only in ("all", "xlsx") else []
    for source in pdfs:
        check_pdf(source, output / "年度报告" / f"{source.stem}.md")
    for source in docxs:
        check_docx(source, output / f"{source.stem}.md")
    for source in xlsxs:
        check_xlsx(source, output / source.stem)
    print(f"\nChecked {CHECKED} source files, {len(ERRORS)} mismatches.")
    return 1 if ERRORS else 0


if __name__ == "__main__":
    sys.exit(main())
