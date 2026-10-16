"""Convert repo inputs (PDF/DOCX/XLSX) to AI-readable text formats.

Layout: originals live under origin/ (unchanged); converted files mirror
them at repo root:
  origin/年度报告/*.pdf -> 年度报告/*.md   (per-page text with page markers)
  origin/*.docx         -> *.md             (headings + paragraphs + markdown tables)
  origin/*.xlsx         -> <name>/*.csv     (cached values per sheet; formula counts logged)

Originals are never modified. Run: uv run --with pdfplumber --with python-docx --with openpyxl python3 scripts/convert_to_text.py
Log: CONVERSION_LOG.md
"""
from __future__ import annotations

import csv
import datetime as _dt
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "origin"
LOG: list[str] = []


def log(msg: str) -> None:
    LOG.append(msg)
    print(msg, flush=True)


def convert_pdf(src: Path, dst: Path) -> None:
    import pdfplumber

    dst.parent.mkdir(parents=True, exist_ok=True)
    empty_pages = 0
    n_pages = 0
    total_chars = 0
    with pdfplumber.open(src) as pdf, dst.open("w", encoding="utf-8") as f:
        f.write(f"# {src.name}\n\n> Source: `{src.as_posix()}` (original unchanged)\n\n")
        for i, page in enumerate(pdf.pages, 1):
            n_pages = i
            text = page.extract_text() or ""
            if not text.strip():
                empty_pages += 1
            total_chars += len(text)
            f.write(f"\n\n---\n\n## Page {i}\n\n{text.strip()}\n")
    warn = f" WARNING: {empty_pages} pages without text layer (scanned?)" if empty_pages else ""
    log(f"PDF {src.name}: {n_pages} pages, {total_chars} chars, {empty_pages} empty.{warn}")


def md_escape_cell(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", "<br>")


def convert_docx(src: Path, dst: Path) -> None:
    from docx import Document

    dst.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(src)
    out: list[str] = [f"# {src.stem}\n\n> Source: `{src.as_posix()}` (original unchanged)\n"]
    n_tables = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower()
        if style.startswith("heading 1"):
            out.append(f"\n# {t}\n")
        elif style.startswith("heading 2"):
            out.append(f"\n## {t}\n")
        elif style.startswith("heading 3"):
            out.append(f"\n### {t}\n")
        else:
            out.append(t)
    for ti, tbl in enumerate(doc.tables, 1):
        n_tables += 1
        out.append(f"\n\n**Table {ti}**\n")
        for ri, row in enumerate(tbl.rows):
            cells = [md_escape_cell(c.text.strip()) for c in row.cells]
            out.append("| " + " | ".join(cells) + " |")
            if ri == 0:
                out.append("| " + " | ".join(["---"] * len(cells)) + " |")
    dst.write_text("\n".join(out) + "\n", encoding="utf-8")
    log(f"DOCX {src.name}: {len(doc.paragraphs)} paras, {n_tables} tables.")


def convert_xlsx(src: Path, outdir: Path) -> None:
    import openpyxl

    wb_vals = openpyxl.load_workbook(src, data_only=True, read_only=True)
    wb_form = openpyxl.load_workbook(src, data_only=False, read_only=True)
    outdir.mkdir(parents=True, exist_ok=True)
    for ws_v in wb_vals.worksheets:
        ws_f = wb_form[ws_v.title]
        safe = re.sub(r"[^\w\-\u4e00-\u9fff]+", "_", ws_v.title).strip("_") or "sheet"
        formula_cells = 0
        rows: list[list[str]] = []
        for row_v, row_f in zip(ws_v.iter_rows(), ws_f.iter_rows()):
            line = []
            for c_v, c_f in zip(row_v, row_f):
                fv = c_f.value
                if isinstance(fv, str) and fv.startswith("="):
                    formula_cells += 1
                v = c_v.value
                line.append("" if v is None else str(v))
            if any(x.strip() for x in line):
                rows.append(line)
        with (outdir / f"{safe}.csv").open("w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rows)
        log(f"XLSX {src.name} [{ws_v.title}]: {len(rows)} non-empty rows -> {safe}.csv"
            + (f" ({formula_cells} formula cells; saved cached values, NOT verified)" if formula_cells else ""))
    wb_vals.close()
    wb_form.close()


def main() -> int:
    if (ROOT / "CONVERSION_LOG.md").exists():
        log("CONVERSION_LOG.md exists; refusing to overwrite. Delete outputs first to re-run.")
        return 1
    pdfs = sorted((SRC / "年度报告").glob("*.pdf"))
    docxs = sorted(SRC.glob("*.docx"))
    xlsxs = sorted(SRC.glob("*.xlsx"))
    log(f"Found {len(pdfs)} PDFs, {len(docxs)} DOCX, {len(xlsxs)} XLSX under origin/.")
    for src in pdfs:
        convert_pdf(src, ROOT / "年度报告" / (src.stem + ".md"))
    for src in docxs:
        convert_docx(src, ROOT / (src.stem + ".md"))
    for src in xlsxs:
        convert_xlsx(src, ROOT / src.stem)
    stamp = _dt.datetime.now(_dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    (ROOT / "CONVERSION_LOG.md").write_text(
        f"# Conversion log ({stamp})\n\nMethod: `scripts/convert_to_text.py` "
        "(pdfplumber page text; python-docx headings/paras/tables; openpyxl cached values per sheet).\n"
        "Originals under `origin/` unchanged. Units/signs/scopes NOT normalized here; verify against filings before quoting.\n\n"
        + "\n".join(f"- {m}" for m in LOG) + "\n",
        encoding="utf-8",
    )
    log("Wrote CONVERSION_LOG.md")
    return 0


if __name__ == "__main__":
    sys.exit(main())
